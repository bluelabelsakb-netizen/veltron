#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Maas ve PuanOzeti sayfalarinin TAM formullerini cikarir."""
import os
import openpyxl

ARSIV = r"C:\Users\pc-n\Documents\Veltron-Excel-arsiv"
XLSX = os.path.join(ARSIV, "Veltron_Is_Takip_Sistemi.xlsx")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "maas-formulleri.txt")

wb = openpyxl.load_workbook(XLSX, data_only=False)
lines = []


def w(t=""):
    lines.append(t)


# --- Maaş sayfası: her sütunun tam formülü (satır 5) ---------------------
ws = wb["Maaş"]
w("=" * 72)
w("MAAS SAYFASI - 5. satir formulleri (sutun sutun)")
w("=" * 72)
for c in range(1, ws.max_column + 1):
    baslik = ws.cell(4, c).value
    form = ws.cell(5, c).value
    if baslik is None and form is None:
        continue
    w(f"\n[{c}] {baslik}")
    w(f"    {form}")

# --- PuanOzeti -----------------------------------------------------------
w("")
w("=" * 72)
w("PUANOZETI - 5. satir formulleri")
w("=" * 72)
ws = wb["PuanOzeti"]
for c in range(1, ws.max_column + 1):
    baslik = ws.cell(4, c).value
    form = ws.cell(5, c).value
    if baslik is None and form is None:
        continue
    w(f"\n[{c}] {baslik}")
    w(f"    {form}")

# --- Puanlar: agirlikli puan formülü -------------------------------------
w("")
w("=" * 72)
w("PUANLAR - 5. satir formulleri")
w("=" * 72)
ws = wb["Puanlar"]
for c in range(1, ws.max_column + 1):
    baslik = ws.cell(4, c).value
    form = ws.cell(5, c).value
    if baslik is None and form is None:
        continue
    w(f"\n[{c}] {baslik}")
    w(f"    {form}")

# --- GenelPerformans -----------------------------------------------------
w("")
w("=" * 72)
w("GENELPERFORMANS - ilk 25 satir")
w("=" * 72)
ws = wb["GenelPerformans"]
for r in range(1, 26):
    vals = []
    for c in range(1, 9):
        v = ws.cell(r, c).value
        if v is None:
            continue
        s = str(v)
        if len(s) > 60:
            s = s[:60] + "..."
        vals.append(f"{c}:{s}")
    if vals:
        w(f"  r{r}: " + " | ".join(vals))

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print(f"Curtiler: {OUT} ({len(lines)} satir)")
